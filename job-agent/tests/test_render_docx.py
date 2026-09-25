from docx import Document

from jobagent.render.docx import build_filename, render_cover_letter, render_resume


def test_build_filename_slugifies_company_and_tag():
    name = build_filename("Resume", "Acme & Co.", "v1 (final)")
    assert name == "Hillary_Mupfumi_Resume_AcmeCo_v1final.docx"


def test_render_resume_creates_readable_docx(tmp_path):
    tailored = {
        "summary": "Data analyst with SQL experience.",
        "bullets": [{"source_id": "b1", "text": "Built SQL dashboards for reporting"}],
        "skills": ["SQL", "Python"],
    }

    out_path = render_resume("Test Candidate", tailored, tmp_path, "Acme", "v1")

    assert out_path.exists()
    assert out_path.name == "Hillary_Mupfumi_Resume_Acme_v1.docx"

    doc = Document(out_path)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Test Candidate" in text
    assert "Built SQL dashboards for reporting" in text
    assert "SQL, Python" in text


def test_render_cover_letter_creates_readable_docx(tmp_path):
    out_path = render_cover_letter("I would be a strong fit.", tmp_path, "Acme", "v1")

    assert out_path.exists()
    assert out_path.name == "Hillary_Mupfumi_CoverLetter_Acme_v1.docx"

    doc = Document(out_path)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "I would be a strong fit." in text
